#!/usr/bin/env python3
"""Gera a política de exemplo usada para demonstrar e testar a skill revisao-politica.

DOCUMENTO DE EXEMPLO, SEM RELAÇÃO COM INSTITUIÇÃO REAL. A "Exemplo Pagamentos IP
S.A." é fictícia. O texto reflete, de propósito, as regras vigentes em março de
2025: as lacunas que a revisão deve apontar vêm só de mudanças normativas
posteriores (Res. BCB 493/2025, 501/2025, 506/2025, 559/2026 e 587/2026). O
gabarito fica fora do plugin, em evals/revisao-politica/ na raiz do repositório.

Uso:
    python gerar_politica_exemplo.py [saida.docx]

Dependência: python-docx (pip install python-docx).
"""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

AVISO = "Documento de exemplo, sem relação com instituição real"

CAPA = [
    ("Exemplo Pagamentos IP S.A. (instituição fictícia)", True),
    ("Política de Prevenção a Fraudes", True),
    ("Versão 3.0", False),
    ("Aprovada pelo Conselho de Administração em 18 de março de 2025", False),
    ("Vigência a partir de 18 de março de 2025", False),
    ("Área responsável: Diretoria de Riscos e Compliance", False),
    ("Próxima revisão ordinária: março de 2026", False),
]

# (número, título, [(número da cláusula, texto)])
SECOES = [
    ("1", "Objetivo", [
        ("1.1", "Esta Política estabelece os princípios, as diretrizes e os procedimentos adotados pela "
                "Exemplo Pagamentos IP S.A. (\"Instituição\") para prevenir, detectar e tratar fraudes na "
                "prestação de serviços de pagamento, em atendimento à Resolução BCB nº 142, de 23 de setembro "
                "de 2021, ao Regulamento anexo à Resolução BCB nº 1, de 12 de agosto de 2020 (Regulamento do "
                "Pix), e à Lei nº 12.865, de 9 de outubro de 2013."),
    ]),
    ("2", "Abrangência", [
        ("2.1", "Aplica-se às contas de pagamento pré-pagas, às transações Pix, aos cartões pré-pagos e às "
                "transferências entre contas mantidas na Instituição."),
        ("2.2", "Aplica-se a administradores, colaboradores, prestadores de serviço e parceiros comerciais "
                "que atuem em nome da Instituição."),
    ]),
    ("3", "Governança e aprovação", [
        ("3.1", "Esta Política é aprovada pelo Conselho de Administração, mediante proposta da Diretoria de "
                "Riscos e Compliance, após manifestação do Comitê de Riscos."),
        ("3.2", "A Diretoria de Riscos e Compliance é a responsável pela implementação desta Política e pela "
                "supervisão das atividades de prevenção a fraudes."),
        ("3.3", "O Comitê de Riscos acompanha mensalmente os indicadores de fraude e as medidas adotadas."),
    ]),
    ("4", "Papéis e responsabilidades", [
        ("4.1", "Conselho de Administração: aprovar esta Política e tomar ciência do relatório mensal de fraudes."),
        ("4.2", "Diretoria Executiva: prover recursos e assegurar a aderência das áreas a esta Política."),
        ("4.3", "Área de Prevenção a Fraudes: operar as regras de monitoramento, analisar alertas, tratar "
                "notificações de infração e propor ajustes de parâmetros."),
        ("4.4", "Área de Operações: executar bloqueios, devoluções e restrições de conta determinados pela "
                "área de Prevenção a Fraudes."),
        ("4.5", "Área Jurídica: acompanhar a regulamentação aplicável e as demandas judiciais relacionadas a fraudes."),
        ("4.6", "Auditoria Interna: avaliar periodicamente a efetividade dos controles desta Política."),
    ]),
    ("5", "Monitoramento de fraudes", [
        ("5.1", "A Instituição utiliza solução de gerenciamento de risco de fraude que contempla as "
                "informações de segurança armazenadas no DICT e identifica transações atípicas ou não "
                "compatíveis com o perfil do cliente."),
        ("5.2", "A Instituição faculta aos clientes o estabelecimento de limites por período e o cadastro "
                "prévio de contas autorizadas a receber valores acima desses limites."),
        ("5.3", "Transações iniciadas por clientes da Instituição com fundada suspeita de fraude podem ser "
                "rejeitadas, observado o limite de tempo diferenciado previsto no Manual de Tempos do Pix."),
        ("5.4", "A avaliação de fraude concentra-se nas transações iniciadas pelos clientes da Instituição, "
                "segundo o perfil do pagador. A conta de destino é considerada apenas pela quantidade de "
                "notificações de infração vinculadas ao recebedor no DICT, como um dos fatores de pontuação."),
    ]),
    ("6", "Pix: autenticação e segurança", [
        ("6.1", "A iniciação de transações Pix e os processos relativos a chaves Pix por clientes pessoa "
                "natural somente são aceitos a partir de dispositivo de acesso previamente cadastrado."),
        ("6.2", "O acesso ao aplicativo exige autenticação com senha e fator adicional (biometria ou token)."),
        ("6.3", "A Instituição disponibiliza no aplicativo informações sobre os cuidados que os clientes "
                "devem ter para evitar fraudes."),
    ]),
    ("7", "Abertura de contas", [
        ("7.1", "A abertura de conta exige validação de documento de identificação, biometria facial com "
                "prova de vida e consulta a bases públicas e privadas de dados cadastrais."),
        ("7.2", "Contas abertas há menos de 30 dias operam com limites reduzidos de recebimento e de envio."),
    ]),
    ("8", "Chaves Pix", [
        ("8.1", "O registro de chave Pix é solicitado ao DICT a pedido do cliente."),
        ("8.2", "A Instituição poderá rejeitar pedido de registro de chave Pix de cliente com notificação de "
                "infração armazenada no DICT, a critério da área de Prevenção a Fraudes, comunicando ao "
                "cliente o motivo da rejeição."),
        ("8.3", "Os pedidos de portabilidade e de reivindicação de posse de chave Pix são processados "
                "conforme o Manual Operacional do DICT."),
    ]),
    ("9", "Bloqueio cautelar", [
        ("9.1", "Havendo suspeita de fraude, os recursos recebidos via Pix são bloqueados cautelarmente, "
                "simultaneamente ao crédito, pelo prazo máximo de 72 horas, com comunicação imediata ao cliente."),
        ("9.2", "Durante o bloqueio, a área de Prevenção a Fraudes avalia os indícios. Confirmada a fundada "
                "suspeita, os recursos são devolvidos pelo Mecanismo Especial de Devolução; caso contrário, "
                "o bloqueio cessa imediatamente."),
        ("9.3", "A possibilidade de bloqueio cautelar consta, em cláusula destacada, do contrato de conta."),
    ]),
    ("10", "Mecanismo Especial de Devolução (MED)", [
        ("10.1", "Recebida reclamação de cliente pagador com fundada suspeita de fraude, a Instituição abre "
                 "notificação de infração e solicitação de devolução no DICT."),
        ("10.2", "Recebida notificação de infração na qualidade de prestadora do recebedor, a Instituição "
                 "bloqueia o saldo disponível na conta do recebedor no momento do recebimento da notificação, "
                 "até o valor solicitado, observados os prazos do Manual Operacional do DICT."),
        ("10.3", "O MED alcança exclusivamente a conta que recebeu diretamente a transação contestada. "
                 "Recursos já transferidos pelo recebedor a outras contas não são objeto de bloqueio ou de "
                 "devolução pela Instituição."),
        ("10.4", "Aceita a notificação de infração, a devolução é efetuada a partir da conta do recebedor, "
                 "nos prazos do Manual Operacional do DICT."),
        ("10.5", "A possibilidade de bloqueios e devoluções, inclusive parciais, consta, em cláusula "
                 "destacada, do contrato de conta."),
    ]),
    ("11", "Restrições a contas notificadas", [
        ("11.1", "Aceita notificação de infração, a área de Operações restringe as transações Pix da conta "
                 "envolvida, ressalvadas as devoluções."),
        ("11.2", "Pedidos de revisão de restrição apresentados pelo cliente são analisados pela área de "
                 "Prevenção a Fraudes, sem prazo predeterminado."),
    ]),
    ("12", "Registros e reporte", [
        ("12.1", "A área de Prevenção a Fraudes mantém registros diários das ocorrências de fraude e de "
                 "tentativas de fraude, com as medidas corretivas adotadas."),
        ("12.2", "Mensalmente é elaborado relatório consolidado das ocorrências e das medidas preventivas e "
                 "corretivas, encaminhado ao Comitê de Riscos, à Auditoria Interna, à Diretoria Executiva e "
                 "ao Conselho de Administração."),
    ]),
    ("13", "Treinamento e conscientização", [
        ("13.1", "Os colaboradores das áreas de atendimento, operações e prevenção a fraudes passam por "
                 "treinamento anual sobre esta Política."),
        ("13.2", "A Instituição mantém campanhas periódicas de conscientização de clientes sobre golpes."),
    ]),
    ("14", "Revisão", [
        ("14.1", "Esta Política é revista anualmente ou sempre que houver alteração relevante da "
                 "regulamentação aplicável."),
    ]),
]


def _aviso(par):
    run = par.add_run(AVISO.upper())
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER


def gerar(caminho: str) -> str:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    _aviso(doc.sections[0].header.paragraphs[0])
    _aviso(doc.sections[0].footer.paragraphs[0])
    _aviso(doc.add_paragraph())

    for texto, negrito in CAPA:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        r.bold = negrito
        if negrito:
            r.font.size = Pt(14)

    for numero, titulo, clausulas in SECOES:
        doc.add_heading(f"{numero}. {titulo}", level=2)
        for num, texto in clausulas:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.add_run(f"{num} ").bold = True
            p.add_run(texto)

    doc.add_paragraph()
    p = doc.add_paragraph("Histórico de versões: 1.0 (nov. 2021); 2.0 (jan. 2023); 3.0 (mar. 2025, "
                          "revisão anual e adequação às Resoluções BCB nº 402/2024, 403/2024 e 425/2024).")
    _aviso(doc.add_paragraph())

    os.makedirs(os.path.dirname(caminho) or ".", exist_ok=True)
    doc.save(caminho)
    return caminho


if __name__ == "__main__":
    destino = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "politica-prevencao-fraudes-exemplo.docx")
    print(gerar(destino))
